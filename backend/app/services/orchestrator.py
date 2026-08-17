"""
Research Orchestrator for ResearchGuard AI.
Coordinates the autonomous multi-agent pipeline in sequence:
1. ResearchPlannerAgent: Sub-questions & search strategy
2. LiteratureSearchAgent: Academic API querying & curation
3. EvidenceExtractionAgent: Quantitative claim & metric extraction
4. CitationVerifierAgent: Rigorous claim grounding & verdict determination
5. ResearchCriticAgent: Adversarial peer-review & bias detection
6. ReportWriterAgent: Publication-grade verified scientific report incorporating uploaded document text
7. Diagram & Presentation Agents: Mindmaps, flowcharts, and slide deck
8. CostTracker: Token usage, latency, and USD cost accounting
Broadcasts real-time WebSocket state changes to frontend.
"""
from datetime import datetime
from bson import ObjectId
from typing import Optional, Dict, Any
import asyncio
import time
import os
import PyPDF2
import docx

from app.core.database import get_db
from app.services.research_planner import ResearchPlannerAgent
from app.services.literature_search import LiteratureSearchAgent
from app.services.evidence_agent import EvidenceExtractionAgent
from app.services.citation_verifier import CitationVerifierAgent
from app.services.critic_agent import ResearchCriticAgent
from app.services.report_writer import ReportWriterAgent
from app.services.agents import DiagramAgent, PresentationAgent
from app.services.cost_tracker import CostTracker
from app.api.routes.websocket import send_agent_update, send_progress, send_complete


def extract_text_from_file(file_path: str) -> str:
    if not os.path.exists(file_path):
        return ""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    if ext == ".pdf":
        try:
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                pages_text = []
                for i, page in enumerate(reader.pages):
                    try:
                        extracted = page.extract_text()
                        if extracted and extracted.strip():
                            pages_text.append(f"--- Page {i+1} ---\n{extracted.strip()}")
                    except Exception:
                        continue
                text = "\n\n".join(pages_text)
        except Exception as e:
            print(f"[PDF_EXTRACT_ERROR] {e}")
            text = ""
        return text.strip()
    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        except Exception:
            return ""
    elif ext in [".txt", ".md", ".csv", ".json"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read().strip()
        except Exception:
            return ""
    return ""


class ResearchOrchestrator:

    @staticmethod
    async def run(
        project_id: str,
        topic: str,
        description: Optional[str] = None,
        url: Optional[str] = None,
        file_path: Optional[str] = None,
        research_mode: str = "literature_review",
    ):
        db = get_db()
        pid = ObjectId(project_id)
        cost_tracker = CostTracker(project_id)

        async def update_agent(agent_name: str, status: str, message: str, data=None):
            await db.research_projects.update_one(
                {"_id": pid},
                {
                    "$set": {
                        f"agents_status.{agent_name}": status,
                        "updated_at": datetime.utcnow(),
                    }
                },
            )
            await db.agent_logs.insert_one({
                "project_id": project_id,
                "agent_name": agent_name,
                "status": status,
                "message": message,
                "data": data,
                "timestamp": datetime.utcnow(),
            })
            await send_agent_update(project_id, agent_name, status, message, data)

        try:
            # ── 1. Pipeline Initialization ────────────────────────
            await db.research_projects.update_one(
                {"_id": pid},
                {
                    "$set": {
                        "status": "processing",
                        "updated_at": datetime.utcnow(),
                        "agents_status": {
                            "coordinator": "running",
                            "planner": "idle",
                            "literature": "idle",
                            "evidence": "idle",
                            "verifier": "idle",
                            "critic": "idle",
                            "report": "idle",
                        }
                    }
                },
            )
            await update_agent("coordinator", "running", "🎯 Lead Coordinator initializing ResearchGuard multi-agent pipeline...")
            await send_progress(project_id, 5, "Initializing pipeline...")
            await asyncio.sleep(0.3)

            # Read uploaded file if provided
            file_content = None
            if file_path and os.path.exists(file_path):
                try:
                    file_content = extract_text_from_file(file_path)
                    await update_agent("coordinator", "running", f"📄 Extracted {len(file_content)} characters from uploaded manuscript/template.")
                except Exception as e:
                    await update_agent("coordinator", "running", f"⚠️ Could not read document: {str(e)}")

            # ── 2. Planner Agent ─────────────────────────────────
            await update_agent("planner", "running", "🎯 Deconstructing research inquiry into academic sub-questions & search criteria...")
            await send_progress(project_id, 15, "Formulating research plan...")
            
            t0 = time.time()
            planner = ResearchPlannerAgent()
            plan_data, plan_usage = await planner.run(
                topic=topic,
                description=description,
                url=url,
                file_content=file_content,
                research_mode=research_mode,
            )
            cost_tracker.record_usage("planner", plan_usage, duration_s=time.time() - t0)

            # Store plan
            await db.research_plans.update_one(
                {"project_id": project_id},
                {"$set": {"project_id": project_id, "plan": plan_data, "created_at": datetime.utcnow()}},
                upsert=True
            )
            sub_count = len(plan_data.get("sub_questions", []))
            await update_agent("planner", "completed", f"✅ Research plan formalized with {sub_count} scientific sub-questions")
            await send_progress(project_id, 25, "Plan generated")

            # ── 3. Literature Search Agent ────────────────────────
            await update_agent("literature", "running", "🔎 Querying academic repositories (arXiv, Crossref, Semantic Scholar)...")
            await send_progress(project_id, 35, "Curating peer-reviewed literature...")

            t0 = time.time()
            lit_agent = LiteratureSearchAgent()
            sources, lit_usage = await lit_agent.run(
                research_plan=plan_data,
                uploaded_doc_text=file_content,
                provided_url=url,
            )
            cost_tracker.record_usage("literature", lit_usage, duration_s=time.time() - t0)

            # Store sources
            await db.literature_sources.update_one(
                {"project_id": project_id},
                {"$set": {"project_id": project_id, "sources": sources, "created_at": datetime.utcnow()}},
                upsert=True
            )
            await update_agent("literature", "completed", f"✅ Curated and quality-scored {len(sources)} peer-reviewed publications")
            await send_progress(project_id, 45, "Literature curated")

            # ── 4. Evidence Extraction Agent ──────────────────────
            await update_agent("evidence", "running", "🧬 Extracting empirical claims, sample sizes, and quantitative effect sizes...")
            await send_progress(project_id, 55, "Extracting empirical evidence...")

            t0 = time.time()
            evidence_agent = EvidenceExtractionAgent()
            evidence_items, ev_usage = await evidence_agent.run(
                research_plan=plan_data,
                sources=sources,
                document_text=file_content,
            )
            cost_tracker.record_usage("evidence", ev_usage, duration_s=time.time() - t0)

            # Store evidence
            await db.evidence_items.update_one(
                {"project_id": project_id},
                {"$set": {"project_id": project_id, "evidence": evidence_items, "created_at": datetime.utcnow()}},
                upsert=True
            )
            await update_agent("evidence", "completed", f"✅ Extracted {len(evidence_items)} testable empirical claims")
            await send_progress(project_id, 65, "Evidence extracted")

            # ── 5. Citation Verifier Agent ────────────────────────
            await update_agent("verifier", "running", "🔗 Cross-examining claims against source texts to audit citation grounding...")
            await send_progress(project_id, 75, "Auditing citation grounding...")

            t0 = time.time()
            verifier = CitationVerifierAgent()
            verifications, ver_usage = await verifier.run(
                evidence_items=evidence_items,
                sources=sources,
            )
            cost_tracker.record_usage("verifier", ver_usage, duration_s=time.time() - t0)

            # Store verifications
            await db.claim_verifications.update_one(
                {"project_id": project_id},
                {"$set": {"project_id": project_id, "verifications": verifications, "created_at": datetime.utcnow()}},
                upsert=True
            )
            supp_count = sum(1 for v in verifications if v.get("verdict") == "SUPPORTED")
            await update_agent("verifier", "completed", f"✅ Citation audit finished: {supp_count}/{len(verifications)} claims verified supported")
            await send_progress(project_id, 80, "Verifications completed")

            # ── 6. Adversarial Critic Agent ───────────────────────
            await update_agent("critic", "running", "🛡️ Performing adversarial peer review, bias checks, and correlation vs causation audit...")
            await send_progress(project_id, 85, "Adversarial stress-testing...")

            t0 = time.time()
            critic = ResearchCriticAgent()
            critiques, crit_usage = await critic.run(
                evidence_items=evidence_items,
                verifications=verifications,
                sources=sources,
            )
            cost_tracker.record_usage("critic", crit_usage, duration_s=time.time() - t0)

            # Store critiques
            await db.critiques.update_one(
                {"project_id": project_id},
                {"$set": {"project_id": project_id, "critiques": critiques, "created_at": datetime.utcnow()}},
                upsert=True
            )
            await update_agent("critic", "completed", f"✅ Adversarial audit logged {len(critiques)} methodological stress-test items")
            await send_progress(project_id, 90, "Critiques completed")

            # ── 7. Report Writer Agent & Output Synthesis ─────────
            await update_agent("report", "running", "📝 Synthesizing publication-grade verified scientific report...")
            await send_progress(project_id, 93, "Writing verified report...")

            t0 = time.time()
            report_writer = ReportWriterAgent()
            report_data, rep_usage = await report_writer.run(
                research_plan=plan_data,
                sources=sources,
                evidence_items=evidence_items,
                verifications=verifications,
                critiques=critiques,
                document_text=file_content,
            )
            cost_tracker.record_usage("report", rep_usage, duration_s=time.time() - t0)

            # Save final report doc
            report_doc = {
                "project_id": project_id,
                "title": report_data.get("title", f"Scientific Evidence Report: {topic}"),
                "executive_summary": report_data.get("executive_summary", ""),
                "findings": report_data.get("findings", []),
                "key_insights": report_data.get("key_insights", []),
                "critic_evaluation": report_data.get("critic_evaluation", ""),
                "recommendations": report_data.get("recommendations", []),
                "references": report_data.get("references", []),
                "integrity_score": report_data.get("integrity_score"),
                "safety_disclaimer": report_data.get("safety_disclaimer", ""),
                "created_at": datetime.utcnow(),
            }
            await db.reports.update_one(
                {"project_id": project_id},
                {"$set": report_doc},
                upsert=True
            )

            # Generate Diagrams & Slides in parallel for rich visualization
            try:
                summary_proxy = {
                    "executive_summary": report_data.get("executive_summary", ""),
                    "key_insights": report_data.get("key_insights", []),
                    "main_themes": [plan_data.get("domain", topic)],
                }
                research_proxy = {
                    "overview": report_data.get("executive_summary", ""),
                    "key_concepts": plan_data.get("sub_questions", [])[:4],
                    "findings": [{"point": f.get("section", ""), "detail": f.get("content", "")[:200]} for f in report_data.get("findings", [])[:5]],
                    "references": report_data.get("references", []),
                }

                diagram_agent = DiagramAgent()
                presentation_agent = PresentationAgent()
                diagram_data, presentation_data = await asyncio.gather(
                    diagram_agent.run(topic, research_proxy, summary_proxy),
                    presentation_agent.run(topic, research_proxy, summary_proxy),
                )

                # Save Diagrams
                await db.diagrams.update_one(
                    {"project_id": project_id},
                    {"$set": {
                        "project_id": project_id,
                        "diagram_type": "mindmap",
                        "mermaid_code": diagram_data.get("mindmap_code", ""),
                        "mindmap_code": diagram_data.get("mindmap_code", ""),
                        "flowchart_code": diagram_data.get("flowchart_code", ""),
                        "title": diagram_data.get("diagram_title", f"{topic} Knowledge Map"),
                        "created_at": datetime.utcnow(),
                    }},
                    upsert=True
                )

                # Save Slides
                await db.slides.update_one(
                    {"project_id": project_id},
                    {"$set": {
                        "project_id": project_id,
                        "title": presentation_data.get("presentation_title", f"{topic} Presentation"),
                        "slides": presentation_data.get("slides", []),
                        "created_at": datetime.utcnow(),
                    }},
                    upsert=True
                )
            except Exception as viz_err:
                print(f"Visualization generation fallback: {viz_err}")

            score_display = f" with Integrity Score {report_data.get('integrity_score')}/100" if report_data.get('integrity_score') is not None else ""
            await update_agent("report", "completed", f"✅ Publication-grade report synthesized{score_display}")

            # ── 8. Cost & Token Accounting ────────────────────────
            verdict_counts = {
                "supported": sum(1 for v in verifications if v.get("verdict") == "SUPPORTED"),
                "partially_supported": sum(1 for v in verifications if v.get("verdict") == "PARTIALLY_SUPPORTED"),
                "contradicted": sum(1 for v in verifications if v.get("verdict") == "CONTRADICTED"),
                "unsupported": sum(1 for v in verifications if v.get("verdict") in ["UNSUPPORTED", "SOURCE_NOT_FOUND"]),
            }
            cost_summary = cost_tracker.get_summary(
                sources_count=len(sources),
                evidence_count=len(evidence_items),
                verifications_count=len(verifications),
                verdicts=verdict_counts,
            )
            await db.cost_summaries.update_one(
                {"project_id": project_id},
                {"$set": cost_summary},
                upsert=True
            )

            # ── 9. Finalize Project ───────────────────────────────
            await db.research_projects.update_one(
                {"_id": pid},
                {
                    "$set": {
                        "status": "completed",
                        "integrity_score": report_data.get("integrity_score"),
                        "has_report": True,
                        "has_slides": True,
                        "has_diagram": True,
                        "updated_at": datetime.utcnow(),
                    }
                },
            )
            await update_agent("coordinator", "completed", "🎉 Verification dossier complete! All agents finished successfully.")
            await send_progress(project_id, 100, "Research complete!")
            await send_complete(project_id)

        except Exception as e:
            print(f"Orchestrator error for {project_id}: {str(e)}")
            await db.research_projects.update_one(
                {"_id": pid},
                {"$set": {"status": "failed", "error": str(e), "updated_at": datetime.utcnow()}},
            )
            await update_agent("coordinator", "failed", f"Pipeline error: {str(e)}")
